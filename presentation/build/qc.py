"""Programmatic QC of the built deck."""
import re, sys, collections, json, pathlib
from pptx import Presentation
from pptx.util import Emu
ROOT = pathlib.Path(__file__).resolve().parents[2]
prs = Presentation(str(ROOT / "presentation" / "CBAS_Interim_Presentation.pptx"))
W, H = prs.slide_width, prs.slide_height
fonts = collections.Counter(); nums = collections.Counter(); problems = []
for i, sl in enumerate(prs.slides, 1):
    notes = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
    if len(notes.split()) < 40: problems.append(f"slide {i}: notes too short ({len(notes.split())} words)")
    alltext = []
    for sh in sl.shapes:
        l, t, r, b = sh.left, sh.top, sh.left + sh.width, sh.top + sh.height
        if l < -Emu(10000) or t < -Emu(10000) or r > W + Emu(10000) or b > H + Emu(10000):
            problems.append(f"slide {i}: shape outside slide: {sh.shape_type} at ({l/914400:.2f},{t/914400:.2f}) size ({sh.width/914400:.2f}x{sh.height/914400:.2f})")
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r_ in p.runs:
                    fonts[r_.font.name] += 1
                    txt = r_.text
                    alltext.append(txt)
                    if any(tok in txt for tok in ("[[", "]]", "{{", "}}", "_{", "^{")) or ("$" in txt and r_.font.name != "Courier New"):
                        problems.append(f"slide {i}: stray markup in run: {txt[:60]!r}")
    for n in re.findall(r"\d+\.\d{2,3}", " ".join(alltext) + " " + notes):
        nums[n] += 1
print("slides:", len(prs.slides))
print("fonts:", dict(fonts))
print("problems:", len(problems))
for p in problems: print("  ", p)
print("decimals used (>=2 places):", sorted(nums.items(), key=lambda kv: float(kv[0])))
