"""Write presentation/speaker_notes.md and print the per-tag slide lists."""
import pathlib, re, collections
from pptx import Presentation
ROOT = pathlib.Path(__file__).resolve().parents[2]
prs = Presentation(str(ROOT / "presentation" / "CBAS_Interim_Presentation.pptx"))
TAGS = {"PAPER", "REPRODUCED", "IMPLEMENTATION", "EXPERIMENT", "OUR AUDIT"}
out = ["# Speaker notes: CBAS interim presentation\n"]
bytag = collections.defaultdict(list); rows = []
for i, sl in enumerate(prs.slides, 1):
    texts = [(sh.top, sh.left, sh.text_frame.text) for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    tags = [t for top, left, t in texts if t in TAGS and top < 600000]
    title = next((t for top, left, t in sorted(texts) if 450000 < top < 700000 and t not in TAGS), "(title or divider slide)")
    title = {1: "Title slide", 32: "Thank you (the talk ends here)", 33: "Appendix divider (backup, only if asked)"}.get(i, title)
    label = next((m.group(1) for _, _, tx in texts if (m := re.fullmatch(r"Backup (B\d+)", tx.strip()))), None)
    num = label if label else str(i)
    for t in tags: bytag[t].append(num)
    rows.append((num, title.replace("\n", " "), tags))
    out.append(f"\n## Slide {num}: {title}\n")
    if tags: out.append(f"Tags: {', '.join(tags)}\n")
    out.append((sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else "") + "\n")
(ROOT / "presentation" / "speaker_notes.md").write_text("\n".join(out))
for t in ["PAPER", "REPRODUCED", "IMPLEMENTATION", "EXPERIMENT", "OUR AUDIT"]:
    print(f"{t:<15} {', '.join(bytag[t])}")
print("\nwords in notes:", sum(len(o.split()) for o in out))
