"""Deck-building primitives: markup -> runs, shapes, arrows, tables, terminals, overflow checks.

Markup used in slide text
    $...$        math: Latin letters italic Cambria, digits/operators upright
    x_{i} x^{k}  sub/superscript (baseline runs; the renderer reduces the size)
    {{...}}      upright (non-italic) text inside math
    [[red|...]]  colour span (ink mut red grn blu vio amb slt wht acc tgrn tred tamb tmut)
    **...**      bold                `...`  monospace
"""
from __future__ import annotations

import pathlib
import sys

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import theme as T  # noqa: E402

COLORS = {
    "ink": T.INK, "mut": T.MUTED, "red": T.BAD, "grn": T.GOOD, "blu": T.SEN, "vio": T.KGC, "amb": T.AGG,
    "slt": T.VER, "wht": "#FFFFFF", "acc": T.GOOD, "tgrn": "#7BD88F", "tred": "#FF8F85", "tamb": "#FFD479",
    "tmut": "#8FA1B3", "ter": T.TERM_FG,
}
PARTY = {  # (line colour, tint)
    "kgc": (T.KGC, T.KGC_T), "sen": (T.SEN, T.SEN_T), "agg": (T.AGG, T.AGG_T), "ver": (T.VER, T.VER_T),
    "bad": (T.BAD, T.BAD_T), "good": (T.GOOD, T.GOOD_T), "plain": (T.RULE, T.WHITE), "panel": (T.RULE, T.PANEL),
}
GREEK_UP = set("ΔΣΠΓΘΛΞΦΨΩ")
WARN: list[str] = []
CUR = {"slide": "?"}

# ------------------------------------------------------------------ colour helpers

def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


def col(c: str) -> str:
    return COLORS.get(c, c)


# ------------------------------------------------------------------ markup parser

class Run:
    __slots__ = ("t", "kind", "it", "bold", "color", "base")

    def __init__(self, t, kind, it, bold, color, base):
        self.t, self.kind, self.it, self.bold, self.color, self.base = t, kind, it, bold, color, base

    def key(self):
        return (self.kind, self.it, self.bold, self.color, self.base)


def _emit(out, text, st):
    if not text:
        return
    kind = "mono" if st["code"] else ("math" if st["math"] else "body")
    color = st["colors"][-1]
    base = {"sub": -25000, "sup": 30000, None: 0}[st["script"]]
    if kind == "math" and not st["up"]:
        buf, cur = "", None
        for ch in text:
            it = ch.isalpha() and ch not in GREEK_UP
            if cur is None or it == cur:
                buf += ch
            else:
                out.append(Run(buf, kind, cur, st["bold"], color, base))
                buf = ch
            cur = it
        out.append(Run(buf, kind, cur, st["bold"], color, base))
    else:
        out.append(Run(text, "math" if kind == "math" else kind, False, st["bold"], color, base))


def _parse(s, st, out):
    i, n = 0, len(s)
    while i < n:
        if s.startswith("[[", i) and "|" in s[i:]:
            j = s.index("|", i)
            st["colors"].append(col(s[i + 2:j]))
            i = j + 1
        elif s.startswith("]]", i):
            if len(st["colors"]) > 1:
                st["colors"].pop()
            i += 2
        elif s[i] == "\\" and i + 1 < n:
            _emit(out, s[i + 1], st)
            i += 2
        elif s[i] == "$":
            st["math"] = not st["math"]
            i += 1
        elif s.startswith("**", i):
            st["bold"] = not st["bold"]
            i += 2
        elif s[i] == "`":
            st["code"] = not st["code"]
            i += 1
        elif s.startswith("{{", i):
            st["up"] = True
            i += 2
        elif s.startswith("}}", i):
            st["up"] = False
            i += 2
        elif s[i] in "_^" and i + 1 < n and s[i + 1] == "{":
            depth, j = 1, i + 2
            while j < n and depth:
                depth += {"{": 1, "}": -1}.get(s[j], 0)
                j += 1
            st2 = dict(st, script="sub" if s[i] == "_" else "sup", colors=list(st["colors"]))
            _parse(s[i + 2:j - 1], st2, out)
            i = j
        else:
            j = i + 1
            while j < n and s[j] not in "[]$`*{}_^\\":
                j += 1
            _emit(out, s[i:j], st)
            i = j


def parse(text: str, color=None, bold=False) -> list[Run]:
    out: list[Run] = []
    _parse(text, {"math": False, "up": False, "bold": bold, "code": False, "colors": [color], "script": None}, out)
    merged: list[Run] = []
    for r in out:
        if merged and merged[-1].key() == r.key():
            merged[-1].t += r.t
        else:
            merged.append(r)
    return merged


# ------------------------------------------------------------------ text measurement

_FP = "/usr/share/fonts/truetype/crosextra/"
_FILES = {("body", False, False): "Carlito-Regular.ttf", ("body", True, False): "Carlito-Bold.ttf",
          ("math", False, False): "Caladea-Regular.ttf", ("math", True, False): "Caladea-Italic.ttf",
          ("mono", False, False): "/usr/share/fonts/truetype/liberation/LiberationMono-Regular.ttf"}
_CACHE: dict = {}


def _width(text, kind, italic, bold, size):
    if kind == "mono":
        path = _FILES[("mono", False, False)]
    elif kind == "math":
        path = _FP + _FILES[("math", italic, False)] if not bold else _FP + "Caladea-Bold.ttf"
    else:
        path = _FP + _FILES[("body", bold, False)]
    key = (path, size)
    if key not in _CACHE:
        _CACHE[key] = ImageFont.truetype(path, max(8, int(size * 10)))
    return _CACHE[key].getlength(text) / 10 / 72


def measure_lines(runs: list[Run], size: float, width: float) -> tuple[int, float]:
    """Greedy wrap estimate. Returns (lines, widest unbreakable token in inches)."""
    tokens: list[tuple[float, bool]] = []  # (width, is_space)
    widest = 0.0
    for r in runs:
        sz = size * (0.62 if r.base else 1.0)
        for part in _split_keep_spaces(r.t):
            w = _width(part, r.kind, r.it, r.bold, sz)
            tokens.append((w, part.isspace()))
    lines, cur, tokw = 1, 0.0, 0.0
    for w, sp in tokens:
        if sp:
            cur += w
            widest = max(widest, tokw)
            tokw = 0.0
            continue
        tokw += w
        if cur + w > width + 1e-6 and cur > 0:
            lines += 1
            cur = w
        else:
            cur += w
    widest = max(widest, tokw)
    return lines, widest


def _split_keep_spaces(t):
    out, buf, sp = [], "", None
    for ch in t:
        is_sp = ch == " "
        if sp is None or is_sp == sp:
            buf += ch
        else:
            out.append(buf)
            buf = ch
        sp = is_sp
    if buf:
        out.append(buf)
    return out


# ------------------------------------------------------------------ text boxes

def _fill_par(p, markup, size, color, bold, align, font, lsp, before, after, spc=None):
    runs = parse(markup, color=None, bold=bold)
    for r in runs:
        run = p.add_run()
        run.text = r.t
        f = run.font
        f.size = Pt(size)
        f.bold = r.bold
        f.italic = r.it
        f.name = {"math": T.F_MATH, "mono": T.F_MONO}.get(r.kind, font or T.F_BODY)
        f.color.rgb = rgb(r.color or color)
        rPr = run._r.get_or_add_rPr()
        if r.base:
            rPr.set("baseline", str(r.base))
        if spc:
            rPr.set("spc", str(spc))
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
    if lsp:
        p.line_spacing = lsp
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    return runs


def _bullet(p, marl=0.28, char="•", color=None):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(marl)))
    pPr.set("indent", str(-Inches(marl)))
    for tag in ("a:buClr", "a:buFont", "a:buChar"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    if color:
        clr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr")).set("val", color.lstrip("#").upper())
    bf = etree.SubElement(pPr, qn("a:buFont"))
    bf.set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def _check(runs_list, sizes, w, h, lsp, label, afters, befores):
    total = 0.0
    for runs, size, aft, bef in zip(runs_list, sizes, afters, befores):
        lines, widest = measure_lines(runs, size, w)
        if widest > w + 0.02:
            WARN.append(f"[{CUR['slide']}] token wider than box ({widest:.2f} > {w:.2f} in): {label}")
        total += lines * size * 1.2 * (lsp or 1.0) / 72 + (aft + bef) / 72
    if total > h + 0.04:
        WARN.append(f"[{CUR['slide']}] text overflows box by {total - h:.2f} in (needs {total:.2f}, has {h:.2f}): {label}")


def tb(s, x, y, w, h, text, size=T.T_BODY, color=T.INK, bold=False, align="l", anchor="t", font=None,
       lsp=None, after=0, before=0, bullets=False, check=True, margin=0.0, spc=None, bullet_color=None):
    """Text box. `text` is a markup string or a list of strings/dicts (text,size,color,bold,align,after,before,bullet)."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    items = text if isinstance(text, list) else text.split("\n")
    rl, sizes, afs, bfs = [], [], [], []
    for i, it in enumerate(items):
        d = it if isinstance(it, dict) else {"text": it}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        sz = d.get("size", size)
        aft = d.get("after", after)
        bef = d.get("before", before)
        runs = _fill_par(p, d["text"], sz, d.get("color", color), d.get("bold", bold), d.get("align", align),
                         font, d.get("lsp", lsp), bef, aft, spc)
        bl = d.get("bullet", bullets)
        if bl:
            _bullet(p, color=bullet_color or d.get("bcolor"))
        wavail = w - (0.28 if bl else 0) - 2 * margin
        rl.append(runs), sizes.append(sz), afs.append(aft), bfs.append(bef)
        if check:
            pass
    if check:
        anyb = bullets or any(isinstance(i, dict) and i.get("bullet") for i in items)
        _check(rl, sizes, w - 2 * margin - (0.28 if anyb else 0), h - 2 * margin, lsp,
               (items[0]["text"] if isinstance(items[0], dict) else items[0])[:48], afs, bfs)
    return box


def bullets(s, x, y, w, h, items, size=T.T_BODY, color=T.INK, after=8, lsp=1.0, bcolor=T.MUTED):
    return tb(s, x, y, w, h, [{"text": t, "bullet": True, "after": after} if isinstance(t, str) else {**t, "bullet": True, "after": t.get("after", after)} for t in items],
              size=size, color=color, lsp=lsp, bullet_color=bcolor)


# ------------------------------------------------------------------ shapes

def box(s, x, y, w, h, fill=None, line=None, lw=1.25, dash=False, radius=0.08, text=None, size=T.T_SMALL, color=T.INK,
        bold=False, align="c", anchor="m", pad=0.08, lsp=None, shape=None, check=True, after=0):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / max(0.01, min(w, h)))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
        if dash:
            sh.line.dash_style = MSO_LINE.DASH
    else:
        sh.line.fill.background()
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if text is not None:
        items = text if isinstance(text, list) else [text]
        rl, sizes, afs, bfs = [], [], [], []
        for i, it in enumerate(items):
            d = it if isinstance(it, dict) else {"text": it}
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            sz = d.get("size", size)
            runs = _fill_par(p, d["text"], sz, d.get("color", color), d.get("bold", bold), d.get("align", align),
                             None, d.get("lsp", lsp), d.get("before", 0), d.get("after", after))
            rl.append(runs), sizes.append(sz), afs.append(d.get("after", after)), bfs.append(d.get("before", 0))
        if check:
            _check(rl, sizes, w - 2 * pad, h - 0.08, lsp, (items[0]["text"] if isinstance(items[0], dict) else items[0])[:48],
                   afs, bfs)
    return sh


def party(s, kind, x, y, w, h, text=None, dash=False, lw=1.5, **kw):
    line, tint = PARTY[kind]
    return box(s, x, y, w, h, fill=tint, line=line, lw=lw, dash=dash, text=text, **kw)


def line(s, x1, y1, x2, y2, color=T.INK, lw=1.75, dash=False, head=False, tail=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(col(color))
    c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = MSO_LINE.DASH
    ln = c.line._get_or_add_ln()
    if tail:
        e = etree.SubElement(ln, qn("a:headEnd"))
        e.set("type", "triangle"), e.set("w", "med"), e.set("len", "med")
    if head:
        e = etree.SubElement(ln, qn("a:tailEnd"))
        e.set("type", "triangle"), e.set("w", "med"), e.set("len", "med")
    return c


def arrow(s, pts, color=T.INK, lw=1.75, dash=False, both=False):
    """Orthogonal polyline with an arrowhead on the last segment (and optionally the first)."""
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        last = i == len(pts) - 2
        line(s, x1, y1, x2, y2, color, lw, dash, head=last, tail=(both and i == 0))


def lock(s, x, y, size=0.26, color=T.BAD):
    bw, bh = size * 0.78, size * 0.52
    bx, by = x + (size - bw) / 2, y + size * 0.46
    body = box(s, bx, by, bw, bh, fill=color, line=None, radius=0.03)
    pts = [(bx + bw * 0.2, by), (bx + bw * 0.2, y + size * 0.22), (bx + bw * 0.34, y + size * 0.06),
           (bx + bw * 0.66, y + size * 0.06), (bx + bw * 0.8, y + size * 0.22), (bx + bw * 0.8, by)]
    fb = s.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(px), Inches(py)) for px, py in pts[1:]], close=False)
    sh = fb.convert_to_shape()
    sh.shadow.inherit = False
    sh.fill.background()
    sh.line.color.rgb = rgb(color)
    sh.line.width = Pt(max(1.4, size * 5.5))
    return body


def chip(s, x, y, w, h, text, kind="plain", size=T.T_SMALL, bold=False, dash=False, color=None, **kw):
    line_c, tint = PARTY[kind]
    return box(s, x, y, w, h, fill=tint, line=line_c, lw=1.2, dash=dash, text=text, size=size, bold=bold,
               color=color or T.INK, radius=min(0.12, h / 2.6), **kw)


def pill(s, x, y, w, h, text, fill, color="#FFFFFF", size=12, bold=True):
    return box(s, x, y, w, h, fill=fill, line=None, radius=h / 2, text=text, size=size, color=color, bold=bold,
               pad=0.04, check=False)


def image(s, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return s.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


# ------------------------------------------------------------------ table-like grid

def grid(s, x, y, colw, rowh, rows, size=T.T_SMALL, header=True, hdr_fill=T.PANEL, hdr_color=T.MUTED, pad=0.1,
         rule=T.RULE, valign="m", hdr_size=None, hdr_spc=60):
    """Draw a table from rectangles. cell = str or dict(text, fill, color, bold, align, size, line)."""
    yy = y
    totw = sum(colw)
    for ri, row in enumerate(rows):
        h = rowh[ri] if isinstance(rowh, list) else rowh
        is_h = header and ri == 0
        if is_h:
            box(s, x, yy, totw, h, fill=hdr_fill, line=None, radius=0)
        xx = x
        for ci, cell in enumerate(row):
            d = cell if isinstance(cell, dict) else {"text": cell}
            if d.get("fill"):
                box(s, xx + 0.02, yy + 0.02, colw[ci] - 0.04, h - 0.04, fill=d["fill"], line=d.get("line"), radius=0.05)
            sz = d.get("size", (hdr_size or size * 0.8) if is_h else size)
            tb(s, xx + pad, yy, colw[ci] - 2 * pad, h, d.get("text", ""), size=sz,
               color=d.get("color", hdr_color if is_h else T.INK), bold=d.get("bold", is_h), align=d.get("align", "l"),
               anchor=valign, spc=(hdr_spc if is_h else None), after=0)
            xx += colw[ci]
        yy += h
        line(s, x, yy, x + totw, yy, rule if not is_h else T.MUTED, 1.0 if not is_h else 1.25)
    return yy


def terminal(s, x, y, w, h, lines, size=14, title=None):
    box(s, x, y, w, h, fill=T.TERM_BG, line="#2B3A4A", lw=1.0, radius=0.1)
    yy = y + 0.14
    if title:
        tb(s, x + 0.2, yy, w - 0.4, 0.3, title, size=12, color="#8FA1B3", font=T.F_MONO)
        yy += 0.34
    tb(s, x + 0.2, yy, w - 0.4, h - (yy - y) - 0.1, [{"text": l, "after": 2} for l in lines], size=size, color=T.TERM_FG,
       font=T.F_MONO)


# ------------------------------------------------------------------ deck + slide chrome

class Deck:
    def __init__(self, total_main: int):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(T.W_IN), Inches(T.H_IN)
        self.total = total_main
        self.count = 0
        self.index: list[dict] = []

    def blank(self):
        return self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def slide(self, eyebrow, title, tags=(), notes="", appendix=None, key=None, title_size=None):
        s = self.blank()
        self.count += 1
        label = appendix if appendix else f"{self.count - (1 if self.count > self.total else 0)}"
        CUR["slide"] = appendix or str(self.count)
        # chrome
        if appendix:
            box(s, 0, 0, T.W_IN, 0.11, fill="#9AA5B1", line=None, radius=0)
        tb(s, 0.6, 0.3, 8.6, 0.28, eyebrow.upper(), size=12, color=T.MUTED, bold=True, spc=90, check=False)
        tb(s, 0.6, 0.58, 12.1, 0.98, title, size=title_size or 30, color=T.INK, bold=True, font=T.F_MATH, anchor="t", lsp=0.95)
        self._tags(s, tags)
        line(s, 0.6, 7.02, 12.73, 7.02, T.RULE, 1.0)
        tb(s, 0.6, 7.08, 9, 0.25, "Certificate-based aggregate signatures for IIoT · interim research presentation",
           size=10.5, color=T.MUTED, check=False)
        tb(s, 10.7, 7.08, 2.03, 0.25, (f"Backup {appendix}" if appendix else f"{CUR['slide']} / {self.total}"),
           size=10.5, color=T.MUTED, align="r", check=False)
        if notes:
            s.notes_slide.notes_text_frame.text = notes.strip()
        self.index.append({"n": CUR["slide"], "title": title, "tags": list(tags), "key": key})
        return s

    def _tags(self, s, tags):
        x = 12.73
        for t in reversed(list(tags)):
            w = 0.16 + 0.092 * len(t)
            x -= w
            pill(s, x, 0.3, w, 0.27, t, T.TAGS[t], size=10.5)
            x -= 0.08

    def dark(self, notes="", appendix=None, bg=None):
        s = self.blank()
        self.count += 1
        CUR["slide"] = appendix or str(self.count)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(bg or T.INK)
        if notes:
            s.notes_slide.notes_text_frame.text = notes.strip()
        self.index.append({"n": CUR["slide"], "title": "(dark)", "tags": [], "key": None})
        return s

    def save(self, path):
        # drop the theme style reference from every shape: it adds a soft shadow in some renderers
        for sl in self.prs.slides:
            for st in list(sl.shapes._spTree.iter(qn("p:style"))):
                st.getparent().remove(st)
        self.prs.save(str(path))


def takeaway(s, text, color=T.INK, y=6.3, size=20, kind=None):
    c = {"red": T.BAD, "grn": T.GOOD}.get(color, color)
    box(s, 0.6, y, 0.09, 0.62, fill=c, line=None, radius=0)
    tb(s, 0.85, y, 11.85, 0.62, text, size=size, bold=True, color=T.INK, anchor="m", lsp=0.95)


def text_width(markup: str, size: float) -> float:
    """Approximate rendered width (inches) of a single-line markup string."""
    tot = 0.0
    for r in parse(markup):
        tot += _width(r.t, r.kind, r.it, r.bold, size * (0.62 if r.base else 1.0))
    return tot


# ------------------------------------------------------------------ presenter cues
CUES = {"on": True}
CUE_FILL, CUE_LINE, CUE_LABEL, CUE_TEXT = "#FFF3B0", "#D9B93A", "#8A6D00", "#4A3B00"


def cue(s, text, x=9.3, y=1.13, w=3.43, h=0.5, name="PresenterCue"):
    """A small sticky-note reminder for the presenter. Not audience content; build with --no-cues to omit."""
    if not CUES["on"]:
        return None
    sh = box(s, x, y, w, h, fill=CUE_FILL, line=CUE_LINE, lw=1.0, radius=0.05,
             text=[{"text": "PRESENTER CUE", "size": 8.5, "bold": True, "color": CUE_LABEL},
                   {"text": text, "size": 12.5, "color": CUE_TEXT}], align="l", pad=0.1, anchor="m", lsp=0.92)
    sh.name = name
    return sh
